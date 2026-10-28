"""Fictional activity imports exercise strict interpretation boundaries."""
import csv
from datetime import date
from pathlib import Path
import tempfile
import unittest

from regnav_esg.tabular import import_activity_rows

try:
    import openpyxl
except ImportError:
    openpyxl = None


HEADERS = ["activity_id", "entity_id", "facility_id", "date", "quantity", "unit", "factor_id", "factor_version", "scope", "scope2_basis", "biogenic"]
ROW = ["fictional-a1", "fictional-entity", "A", "2025-06-30", "2.5", "MWh", "fictional-factor", "v1", "2", "location_based", "false"]


class TabularTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "activities.csv"

    def write(self, rows, delimiter=","):
        with self.path.open("w", encoding="utf-8", newline="") as handle:
            csv.writer(handle, delimiter=delimiter).writerows(rows)
        return self.path

    def test_csv_import_has_exact_quantities_and_typed_flags(self):
        self.write([HEADERS, ROW])
        result = import_activity_rows(self.path)
        self.assertEqual(result[0]["quantity"], "2.5")
        self.assertEqual(result[0]["scope"], 2)
        self.assertIs(result[0]["biogenic"], False)

    def test_explicit_mapping_and_delimiter(self):
        headers = ["Column " + str(index) for index in range(len(HEADERS))] + ["deliberately excluded"]
        mapping = dict(zip(HEADERS, headers))
        self.write([headers, ROW + ["extra information"]], delimiter=";")
        self.assertEqual(import_activity_rows(self.path, mapping=mapping, delimiter=";")[0]["quantity"], "2.5")
        with self.assertRaises(ValueError):
            import_activity_rows(self.path, delimiter=";")

    def test_duplicate_headers_and_ids_rejected(self):
        self.write([HEADERS + ["quantity"], ROW + ["1"]])
        with self.assertRaisesRegex(ValueError, "Duplicate column"):
            import_activity_rows(self.path)
        self.write([HEADERS, ROW, ROW])
        with self.assertRaisesRegex(ValueError, "duplicate activity_id"):
            import_activity_rows(self.path)

    def test_missing_required_mapping_rejected(self):
        self.write([HEADERS, ROW])
        mapping = {field: field for field in HEADERS if field != "factor_version"}
        with self.assertRaisesRegex(ValueError, "Missing required"):
            import_activity_rows(self.path, mapping=mapping)

    def test_locale_and_exponent_quantities_are_not_guessed(self):
        for value in ("1,250", "1.250,00", "1e3", " 1.2", "NaN"):
            row = ROW.copy()
            row[4] = value
            self.write([HEADERS, row])
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "decimal dot"):
                import_activity_rows(self.path)

    def test_dates_are_not_guessed(self):
        for value in ("01/02/2025", "20250630", "2025-02-30", "2025-6-3"):
            row = ROW.copy()
            row[3] = value
            self.write([HEADERS, row])
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "YYYY-MM-DD"):
                import_activity_rows(self.path)

    def test_boolean_is_not_truthy_text(self):
        row = ROW.copy()
        row[10] = "yes"
        self.write([HEADERS, row])
        with self.assertRaisesRegex(ValueError, "true or false"):
            import_activity_rows(self.path)

    def test_formula_text_rejected(self):
        for value in ("=1+1", "\t=HYPERLINK('x')", "+SUM(1)", "@SUM(1)", "-SUM(1)"):
            row = ROW.copy()
            row[0] = value
            self.write([HEADERS, row])
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "Formula"):
                import_activity_rows(self.path)

    def test_rows_may_not_drop_or_gain_cells(self):
        for row in (ROW[:-1], ROW + ["extra"]):
            self.write([HEADERS, row])
            with self.assertRaisesRegex(ValueError, "cell count"):
                import_activity_rows(self.path)

    def test_empty_table_cannot_be_a_zero_inventory(self):
        self.write([HEADERS, []])
        with self.assertRaisesRegex(ValueError, "no activity records"):
            import_activity_rows(self.path)

    def test_zero_is_explicitly_retained(self):
        row = ROW.copy()
        row[4] = "0"
        self.write([HEADERS, row])
        self.assertEqual(import_activity_rows(self.path)[0]["quantity"], "0")

    def test_missing_scope_classification_fails(self):
        row = ROW.copy()
        row[9] = ""
        self.write([HEADERS, row])
        with self.assertRaisesRegex(ValueError, "scope2_basis"):
            import_activity_rows(self.path)

    @unittest.skipIf(openpyxl is None, "optional openpyxl is not installed")
    def test_xlsx_text_quantities_and_explicit_sheet(self):
        path = Path(self.temp.name) / "activities.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.title = "Activities"
        workbook.active.append(HEADERS)
        workbook.active.append(ROW)
        workbook.create_sheet("Other")
        workbook.save(path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "Specify sheet"):
            import_activity_rows(path)
        result = import_activity_rows(path, sheet="Activities")
        self.assertEqual(result[0]["quantity"], "2.5")

    @unittest.skipIf(openpyxl is None, "optional openpyxl is not installed")
    def test_xlsx_rejects_formula_float_and_native_date(self):
        path = Path(self.temp.name) / "activities.xlsx"
        for column, value, message in ((5, "=1+1", "Formula"), (5, 1.25, "decimal"), (4, date(2025, 6, 30), "dates")):
            workbook = openpyxl.Workbook()
            workbook.active.append(HEADERS)
            workbook.active.append(ROW)
            workbook.active.cell(row=2, column=column, value=value)
            workbook.save(path)
            workbook.close()
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, message):
                import_activity_rows(path)


if __name__ == "__main__":
    unittest.main()
