"""Strict, local activity import from CSV or optional XLSX workbooks.

The mapping is {activity_field: exact_column_header}. CSV is UTF-8 (an initial
BOM is accepted). Delimiters are explicit. Decimal quantities use dot notation
without grouping or exponent notation. Spreadsheet decimal quantities and dates
must be stored as text; numeric integers are accepted for quantities and scopes.
"""

import csv
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile


_REQUIRED = ("activity_id", "entity_id", "facility_id", "date", "quantity", "unit",
             "factor_id", "factor_version", "scope")
_OPTIONAL = ("scope2_basis", "scope3_category", "biogenic", "equity_fraction")
_FIELDS = set(_REQUIRED + _OPTIONAL)
_DECIMAL = re.compile(r"[0-9]+(?:\.[0-9]+)?\Z")
_MAX_BYTES = 25 * 1024 * 1024
_MAX_EXPANDED_BYTES = 100 * 1024 * 1024
_MAX_ROWS = 100000
_MAX_COLUMNS = 256


def _header_map(headers, mapping):
    if not headers:
        raise ValueError("The table has no header row")
    if len(headers) > _MAX_COLUMNS:
        raise ValueError("Table exceeds the 256-column import limit")
    if any(not isinstance(header, str) or not header.strip() for header in headers):
        raise ValueError("Every column must have a non-empty text header")
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate column headers are not permitted")
    if mapping is None:
        if set(headers) - _FIELDS:
            raise ValueError("Unknown columns require an explicit field-to-header mapping")
        mapping = {header: header for header in headers}
    if not isinstance(mapping, dict) or not mapping:
        raise ValueError("mapping must be a non-empty field-to-header object")
    if set(mapping) - _FIELDS:
        raise ValueError("mapping contains unsupported activity fields")
    if any(not isinstance(header, str) or not header for header in mapping.values()):
        raise ValueError("Mapped column headers must be strings")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("One column cannot map to multiple activity fields")
    missing = set(_REQUIRED) - set(mapping)
    if missing:
        raise ValueError("Missing required activity mappings: " + ", ".join(sorted(missing)))
    missing_headers = set(mapping.values()) - set(headers)
    if missing_headers:
        raise ValueError("Mapped headers not found: " + ", ".join(sorted(missing_headers)))
    return {field: headers.index(header) for field, header in mapping.items()}


def _formula(value, row, column):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "@", "-")):
        raise ValueError(f"Formula-like cell rejected at row {row}, column {column}")


def _integer(value, field, row, lower, upper):
    if isinstance(value, bool):
        raise ValueError(f"Row {row}: {field} must be an integer")
    if isinstance(value, int):
        result = value
    elif isinstance(value, str) and re.fullmatch(r"[0-9]+", value):
        result = int(value)
    else:
        raise ValueError(f"Row {row}: {field} must be an integer without formatting")
    if not lower <= result <= upper:
        raise ValueError(f"Row {row}: {field} is outside its permitted range")
    return result


def _quantity(value, field, row):
    if isinstance(value, bool) or not isinstance(value, (str, int)):
        raise ValueError(f"Row {row}: {field} requires decimal text or an integer; store spreadsheet decimals as text")
    result = str(value)
    if not _DECIMAL.fullmatch(result):
        raise ValueError(f"Row {row}: {field} requires non-negative decimal dot notation without grouping or exponents")
    if len(result.replace(".", "")) > 50:
        raise ValueError(f"Row {row}: {field} exceeds supported numeric precision")
    return result


def _record(values, positions, row):
    result = {}
    for field, index in positions.items():
        value = values[index]
        if value is None or value == "":
            if field in _REQUIRED:
                raise ValueError(f"Row {row}: missing required {field}")
            continue
        if field in ("quantity", "equity_fraction"):
            result[field] = _quantity(value, field, row)
            if field == "equity_fraction" and Decimal(result[field]) > 1:
                raise ValueError(f"Row {row}: equity_fraction must be between 0 and 1")
        elif field == "scope":
            result[field] = _integer(value, field, row, 1, 3)
        elif field == "scope3_category":
            result[field] = _integer(value, field, row, 1, 15)
        elif field == "biogenic":
            if isinstance(value, bool):
                result[field] = value
            elif value in ("true", "false"):
                result[field] = value == "true"
            else:
                raise ValueError(f"Row {row}: biogenic must be true or false")
        elif field == "date":
            if not isinstance(value, str):
                raise ValueError(f"Row {row}: date must be ISO YYYY-MM-DD text")
            try:
                parsed = date.fromisoformat(value)
            except ValueError as exc:
                raise ValueError(f"Row {row}: date must be ISO YYYY-MM-DD text") from exc
            if parsed.isoformat() != value:
                raise ValueError(f"Row {row}: date must be ISO YYYY-MM-DD text")
            result[field] = value
        else:
            if not isinstance(value, str) or not value.strip():
                raise ValueError(f"Row {row}: {field} must be non-empty text")
            result[field] = value
    if result["scope"] == 2 and not result.get("biogenic", False):
        if result.get("scope2_basis") not in ("location_based", "market_based"):
            raise ValueError(f"Row {row}: Scope 2 requires an explicit scope2_basis")
    if result["scope"] == 3 and not result.get("biogenic", False):
        if "scope3_category" not in result:
            raise ValueError(f"Row {row}: Scope 3 requires an explicit category")
    return result


def _convert(rows, mapping):
    try:
        headers = list(next(rows))
    except StopIteration as exc:
        raise ValueError("The table is empty") from exc
    for index, value in enumerate(headers, 1):
        _formula(value, 1, index)
    positions = _header_map(headers, mapping)
    records, identities = [], set()
    for row_number, values in enumerate(rows, 2):
        if row_number > _MAX_ROWS + 1:
            raise ValueError("Table exceeds the 100,000-row import limit")
        values = list(values)
        if not values or all(value is None or value == "" for value in values):
            continue  # Completely empty formatting rows contain no activity.
        if len(values) != len(headers):
            raise ValueError(f"Row {row_number}: cell count does not match the header")
        for index, value in enumerate(values, 1):
            _formula(value, row_number, index)
        record = _record(values, positions, row_number)
        if record["activity_id"] in identities:
            raise ValueError(f"Row {row_number}: duplicate activity_id {record['activity_id']!r}")
        identities.add(record["activity_id"])
        records.append(record)
    if not records:
        raise ValueError("The table contains no activity records; absence is not a measured zero")
    return records


def import_activity_rows(path, mapping=None, sheet=None, delimiter=","):
    """Return activity dictionaries suitable for calculate_inventory.

    CSV supports a caller-selected one-character delimiter. XLSX requires the
    optional openpyxl package, rejects formulas and ambiguous numeric/date cells,
    and requires a sheet name when a workbook has multiple sheets. Unmapped
    columns may be deliberately excluded only by supplying an explicit mapping.
    """
    if not isinstance(delimiter, str) or len(delimiter) != 1 or delimiter in "\r\n\0":
        raise ValueError("delimiter must be one non-newline character")
    source = Path(path)
    if not source.is_file():
        raise ValueError("Activity input must be an existing local file")
    if source.stat().st_size > _MAX_BYTES:
        raise ValueError("Activity input exceeds the 25 MiB import limit")
    suffix = source.suffix.lower()
    if suffix in (".csv", ".tsv"):
        if sheet is not None:
            raise ValueError("sheet applies only to XLSX input")
        try:
            with source.open("r", encoding="utf-8-sig", newline="") as handle:
                return _convert(iter(csv.reader(handle, delimiter=delimiter, strict=True)), mapping)
        except (UnicodeError, csv.Error) as exc:
            raise ValueError(f"Invalid UTF-8 CSV table: {exc}") from exc
    if suffix != ".xlsx":
        raise ValueError("Supported activity formats are .csv, .tsv and .xlsx")
    if delimiter != ",":
        raise ValueError("delimiter applies only to CSV/TSV input")
    try:
        with ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > _MAX_EXPANDED_BYTES:
                raise ValueError("XLSX expanded content exceeds the 100 MiB import limit")
    except BadZipFile as exc:
        raise ValueError("Invalid XLSX archive") from exc
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX import requires the optional openpyxl dependency") from exc
    try:
        workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("The XLSX workbook could not be read") from exc
    try:
        if sheet is None:
            if len(workbook.sheetnames) != 1:
                raise ValueError("Specify sheet explicitly for a workbook with multiple sheets")
            selected = workbook[workbook.sheetnames[0]]
        else:
            if not isinstance(sheet, str) or sheet not in workbook.sheetnames:
                raise ValueError("The named sheet does not exist")
            selected = workbook[sheet]
        if selected.max_row is not None and selected.max_row > _MAX_ROWS + 1:
            raise ValueError("XLSX sheet exceeds the 100,000-row import limit")
        if selected.max_column is not None and selected.max_column > _MAX_COLUMNS:
            raise ValueError("XLSX sheet exceeds the 256-column import limit")

        def rows():
            for row_number, cells in enumerate(selected.iter_rows(), 1):
                values = []
                for index, cell in enumerate(cells, 1):
                    if cell.data_type == "f":
                        raise ValueError(f"Formula cell rejected at row {row_number}, column {index}")
                    if cell.data_type == "e":
                        raise ValueError(f"Spreadsheet error cell at row {row_number}, column {index}")
                    if isinstance(cell.value, (date, datetime)):
                        raise ValueError(f"Row {row_number}: store spreadsheet dates as ISO YYYY-MM-DD text")
                    values.append(cell.value)
                yield values
        return _convert(iter(rows()), mapping)
    finally:
        workbook.close()
